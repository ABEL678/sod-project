import tempfile
from io import BytesIO

from django.contrib.auth import get_user_model
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse
from openpyxl import load_workbook

from storage.domain.models import Company, MailRule, Project, Section, UserProfile
from storage.domain.services import upload_document


User = get_user_model()


@override_settings(MEDIA_ROOT=tempfile.mkdtemp())
class ViewTests(TestCase):
    def setUp(self):
        self.company = Company.objects.create(name='Компания 1')
        self.other_company = Company.objects.create(name='Компания 2')

        self.user_gip = User.objects.create_user('gip', password='pass')
        UserProfile.objects.create(user=self.user_gip, role='gip', company=self.company)

        self.user_viewer = User.objects.create_user('viewer', password='pass')
        UserProfile.objects.create(user=self.user_viewer, role='viewer', company=self.company)

        self.user_contractor = User.objects.create_user('contractor', password='pass')
        UserProfile.objects.create(user=self.user_contractor, role='contractor', company=self.company)

        self.user_foreign_gip = User.objects.create_user('foreign_gip', password='pass')
        UserProfile.objects.create(
            user=self.user_foreign_gip,
            role='gip',
            company=self.other_company,
        )

        self.project = Project.objects.create(
            company=self.company,
            code='P-1',
            name='Проект',
            status='active',
            gip=self.user_gip,
            created_by=self.user_gip,
        )
        self.project_other = Project.objects.create(
            company=self.company,
            code='P-2',
            name='Другой проект',
            status='active',
            gip=self.user_gip,
            created_by=self.user_gip,
        )
        self.foreign_project = Project.objects.create(
            company=self.other_company,
            code='P-FOREIGN',
            name='Чужой проект',
            status='active',
            gip=self.user_foreign_gip,
            created_by=self.user_foreign_gip,
        )
        self.section = Section.objects.create(
            project=self.project,
            code='АР',
            name='Архитектурные решения',
            order=1,
        )
        self.section_other = Section.objects.create(
            project=self.project_other,
            code='КР',
            name='Конструктивные решения',
            order=1,
        )
        self.foreign_section = Section.objects.create(
            project=self.foreign_project,
            code='ОВ',
            name='Отопление и вентиляция',
            order=1,
        )

    def test_project_list(self):
        self.client.force_login(self.user_gip)
        response = self.client.get(reverse('storage:project_list'))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'storage/project_list.html')
        self.assertContains(response, 'P-1')

    def test_project_list_excludes_other_company(self):
        self.client.force_login(self.user_gip)
        response = self.client.get(reverse('storage:project_list'))
        self.assertNotContains(response, 'P-FOREIGN')

    def test_project_detail(self):
        self.client.force_login(self.user_gip)
        response = self.client.get(reverse('storage:project_detail', args=[self.project.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'storage/project_detail.html')
        self.assertContains(response, 'АР')
        self.assertContains(response, 'Перетащите файлы сюда или нажмите для выбора')
        self.assertContains(response, 'upload-ajax')

    def test_project_detail_foreign_company_404(self):
        self.client.force_login(self.user_gip)
        response = self.client.get(reverse('storage:project_detail', args=[self.foreign_project.pk]))
        self.assertEqual(response.status_code, 404)

    def test_project_detail_contractor_sees_only_allowed_documents(self):
        self._make_contractor_docs()
        self.client.force_login(self.user_contractor)
        response = self.client.get(reverse('storage:project_detail', args=[self.project.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'АР-01')
        self.assertNotContains(response, 'КР-99')

    def test_upload_ajax_success(self):
        self.client.force_login(self.user_gip)
        response = self.client.post(
            reverse('storage:document_upload_ajax', args=[self.project.pk]),
            {
                'section_id': self.section.pk,
                'code': 'АР-01',
                'name': 'План 1 этажа',
                'file': SimpleUploadedFile('plan.pdf', b'data', content_type='application/pdf'),
            },
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'status': 'ok', 'version': 1})
        self.assertEqual(self.project.documents.count(), 1)
        self.assertEqual(self.project.documents.first().current_version.version_number, 1)

    def test_upload_ajax_second_version(self):
        self.client.force_login(self.user_gip)
        url = reverse('storage:document_upload_ajax', args=[self.project.pk])
        payload = {
            'section_id': self.section.pk,
            'code': 'АР-01',
            'name': 'План 1 этажа',
        }
        self.client.post(
            url,
            {**payload, 'file': SimpleUploadedFile('plan.pdf', b'data', content_type='application/pdf')},
        )
        response = self.client.post(
            url,
            {**payload, 'file': SimpleUploadedFile('plan2.pdf', b'data 2', content_type='application/pdf')},
        )
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {'status': 'ok', 'version': 2})

    def test_upload_ajax_permission_denied(self):
        self.client.force_login(self.user_viewer)
        response = self.client.post(
            reverse('storage:document_upload_ajax', args=[self.project.pk]),
            {
                'section_id': self.section.pk,
                'code': 'АР-02',
                'name': 'План',
                'file': SimpleUploadedFile('plan.pdf', b'data', content_type='application/pdf'),
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Нет прав на запись', response.json()['error'])
        self.assertEqual(self.project.documents.count(), 0)

    def test_upload_ajax_wrong_project(self):
        self.client.force_login(self.user_gip)
        response = self.client.post(
            reverse('storage:document_upload_ajax', args=[self.project.pk]),
            {
                'section_id': self.section_other.pk,
                'code': 'КР-01',
                'name': 'Разрез',
                'file': SimpleUploadedFile('plan.pdf', b'data', content_type='application/pdf'),
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Раздел не относится к указанному проекту', response.json()['error'])

    def test_upload_ajax_foreign_project_404(self):
        self.client.force_login(self.user_gip)
        response = self.client.post(
            reverse('storage:document_upload_ajax', args=[self.foreign_project.pk]),
            {
                'section_id': self.foreign_section.pk,
                'code': 'ОВ-01',
                'name': 'Схема',
                'file': SimpleUploadedFile('plan.pdf', b'data', content_type='application/pdf'),
            },
        )
        self.assertEqual(response.status_code, 404)
        self.assertEqual(response.json()['error'], 'Проект не найден')

    def test_upload_ajax_missing_section(self):
        self.client.force_login(self.user_gip)
        response = self.client.post(
            reverse('storage:document_upload_ajax', args=[self.project.pk]),
            {
                'code': 'АР-03',
                'name': 'План',
                'file': SimpleUploadedFile('plan.pdf', b'data', content_type='application/pdf'),
            },
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn('Не выбран раздел', response.json()['error'])

    def _upload(self, code='АР-01', name='План 1 этажа', content=b'data', section=None):
        return upload_document(
            user=self.user_gip,
            project=self.project,
            section=section or self.section,
            code=code,
            name=name,
            doc_type='drawing',
            file=SimpleUploadedFile(f'{code}.pdf', content, content_type='application/pdf'),
        )

    def _upload_foreign(self):
        return upload_document(
            user=self.user_foreign_gip,
            project=self.foreign_project,
            section=self.foreign_section,
            code='ОВ-01',
            name='Схема вентиляции',
            doc_type='drawing',
            file=SimpleUploadedFile('foreign.pdf', b'foreign', content_type='application/pdf'),
        )

    def _make_contractor_docs(self):
        self.user_contractor.userprofile.allowed_sections.add(self.section)
        allowed = self._upload(
            code='АР-01',
            name='Разрешённый чертеж',
            content=b'allowed',
        )
        forbidden_section = Section.objects.create(
            project=self.project,
            code='КР',
            name='Конструктивные решения',
            order=2,
        )
        other = self._upload(
            code='КР-99',
            name='Запрещённый чертеж',
            content=b'other',
            section=forbidden_section,
        )
        return allowed, other

    def test_download_version(self):
        self.client.force_login(self.user_gip)
        version = self._upload(content=b'pdf-body')
        response = self.client.get(reverse('storage:document_download', args=[version.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertEqual(b''.join(response.streaming_content), b'pdf-body')

    def test_download_anonymous_redirects_to_login(self):
        version = self._upload()
        response = self.client.get(reverse('storage:document_download', args=[version.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse('login'), response['Location'])

    def test_download_foreign_company_404(self):
        version = self._upload_foreign()
        self.client.force_login(self.user_gip)
        response = self.client.get(reverse('storage:document_download', args=[version.pk]))
        self.assertEqual(response.status_code, 404)

    def test_download_without_read_permission_forbidden(self):
        version = self._upload()
        self.client.force_login(self.user_contractor)
        response = self.client.get(reverse('storage:document_download', args=[version.pk]))
        self.assertEqual(response.status_code, 403)

    def test_export_vedomost(self):
        self.client.force_login(self.user_gip)
        version = self._upload()
        response = self.client.get(reverse('storage:export_vedomost', args=[self.project.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertIn('attachment', response['Content-Disposition'])
        self.assertIn('vedomost.xlsx', response['Content-Disposition'])

        data = b''.join(response.streaming_content)
        workbook = load_workbook(BytesIO(data))
        sheet = workbook.active
        self.assertEqual(sheet.cell(row=1, column=1).value, 'Шифр')
        self.assertEqual(sheet.cell(row=2, column=1).value, 'АР-01')
        self.assertEqual(sheet.cell(row=2, column=2).value, 'План 1 этажа')
        self.assertIsNotNone(sheet.cell(row=2, column=2).hyperlink)
        target = sheet.cell(row=2, column=2).hyperlink.target
        self.assertTrue(target.startswith('http://') or target.startswith('https://'))
        self.assertIn(f'/download/{version.pk}/', target)

    def test_export_vedomost_foreign_project_404(self):
        self.client.force_login(self.user_gip)
        response = self.client.get(reverse('storage:export_vedomost', args=[self.foreign_project.pk]))
        self.assertEqual(response.status_code, 404)

    def test_export_vedomost_contractor_sees_only_allowed_sections(self):
        self._make_contractor_docs()
        self.client.force_login(self.user_contractor)
        response = self.client.get(reverse('storage:export_vedomost', args=[self.project.pk]))
        self.assertEqual(response.status_code, 200)

        data = b''.join(response.streaming_content)
        sheet = load_workbook(BytesIO(data)).active
        codes = [
            sheet.cell(row=row, column=1).value
            for row in range(2, sheet.max_row + 1)
        ]
        self.assertEqual(codes, ['АР-01'])

    def test_export_anonymous_redirects_to_login(self):
        response = self.client.get(reverse('storage:export_vedomost', args=[self.project.pk]))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse('login'), response['Location'])

    def test_document_detail(self):
        self.client.force_login(self.user_gip)
        version = self._upload()
        response = self.client.get(reverse('storage:document_detail', args=[version.document_id]))
        self.assertEqual(response.status_code, 200)
        self.assertTemplateUsed(response, 'storage/document_detail.html')
        self.assertContains(response, 'История версий')
        self.assertContains(response, 'Замечания')
        self.assertContains(response, 'Скачать')

    def test_document_detail_foreign_company_404(self):
        version = self._upload_foreign()
        self.client.force_login(self.user_gip)
        response = self.client.get(reverse('storage:document_detail', args=[version.document_id]))
        self.assertEqual(response.status_code, 404)

    def test_document_detail_contractor_allowed_section(self):
        allowed, _ = self._make_contractor_docs()
        self.client.force_login(self.user_contractor)
        response = self.client.get(reverse('storage:document_detail', args=[allowed.document_id]))
        self.assertEqual(response.status_code, 200)

    def test_document_detail_contractor_foreign_section_404(self):
        _, other = self._make_contractor_docs()
        self.client.force_login(self.user_contractor)
        response = self.client.get(reverse('storage:document_detail', args=[other.document_id]))
        self.assertEqual(response.status_code, 404)

    def test_document_detail_forbidden_without_rights(self):
        version = self._upload()
        self.client.force_login(self.user_contractor)
        response = self.client.get(reverse('storage:document_detail', args=[version.document_id]))
        self.assertEqual(response.status_code, 404)

    def test_project_detail_links_to_document_detail(self):
        self.client.force_login(self.user_gip)
        version = self._upload()
        response = self.client.get(reverse('storage:project_detail', args=[self.project.pk]))
        self.assertContains(
            response,
            reverse('storage:document_detail', args=[version.document_id]),
        )

    def test_pages_require_login(self):
        for url in (
            reverse('storage:project_list'),
            reverse('storage:project_detail', args=[self.project.pk]),
            reverse('storage:export_vedomost', args=[self.project.pk]),
            reverse('storage:simulate_mail_import', args=[self.project.pk]),
        ):
            response = self.client.get(url)
            self.assertEqual(response.status_code, 302, url)
            self.assertIn(reverse('login'), response['Location'], url)

    def test_login_page_renders(self):
        response = self.client.get(reverse('login'))
        self.assertEqual(response.status_code, 200)
        self.assertContains(response, 'Вход в систему')
        self.assertContains(response, 'Войти')

    def test_logout_redirects_home(self):
        self.client.force_login(self.user_gip)
        response = self.client.post(reverse('logout'))
        self.assertEqual(response.status_code, 302)
        self.assertEqual(response['Location'], '/')

    def test_upload_ajax_requires_login(self):
        response = self.client.post(
            reverse('storage:document_upload_ajax', args=[self.project.pk]),
            {
                'section_id': self.section.pk,
                'code': 'АР-99',
                'name': 'План',
                'file': SimpleUploadedFile('plan.pdf', b'data', content_type='application/pdf'),
            },
        )
        self.assertEqual(response.status_code, 403)
        self.assertIn('Требуется авторизация', response.json()['error'])

    def test_simulate_mail_import_success(self):
        MailRule.objects.create(
            project=self.project,
            section=self.section,
            email_filter='P-1',
        )
        self.client.force_login(self.user_gip)
        response = self.client.post(
            reverse('storage:simulate_mail_import', args=[self.project.pk]),
            follow=True,
        )
        self.assertRedirects(response, reverse('storage:project_detail', args=[self.project.pk]))
        self.assertContains(response, 'Письмо обработано, файл добавлен.')
        document = self.project.documents.get()
        self.assertEqual(document.doc_type, 'source_data')
        self.assertEqual(document.current_version.version_number, 1)

    def test_simulate_mail_import_foreign_project_404(self):
        MailRule.objects.create(
            project=self.foreign_project,
            section=self.foreign_section,
            email_filter='P-FOREIGN',
        )
        self.client.force_login(self.user_gip)
        response = self.client.post(
            reverse('storage:simulate_mail_import', args=[self.foreign_project.pk]),
        )
        self.assertEqual(response.status_code, 404)

    def test_simulate_mail_import_without_rules(self):
        self.client.force_login(self.user_gip)
        response = self.client.post(
            reverse('storage:simulate_mail_import', args=[self.project.pk]),
            follow=True,
        )
        self.assertContains(response, 'не настроены правила импорта')
        self.assertEqual(self.project.documents.count(), 0)
