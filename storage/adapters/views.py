import json
import uuid
import zipfile
from io import BytesIO

from django.contrib import messages
from django.contrib.auth import get_user_model
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied
from django.core.files.uploadedfile import SimpleUploadedFile
from django.db.models import Q
from django.http import (
    FileResponse,
    Http404,
    HttpResponse,
    HttpResponseBadRequest,
    HttpResponseForbidden,
    HttpResponseNotFound,
    JsonResponse,
)
from django.shortcuts import get_object_or_404, redirect
from django.urls import reverse
from django.utils import timezone
from django.utils.safestring import mark_safe
from django.utils.text import get_valid_filename
from django.views import View
from django.views.generic import DetailView, ListView, TemplateView
from django_scotty.views import CottonTableView
from openpyxl import Workbook
from openpyxl.styles import Font

from storage.adapters.filters import DocumentFilter
from storage.adapters.local_storage import LocalFileStorage
from storage.adapters.tables import DocumentTable
from storage.domain.models import (
    Company,
    Document,
    DocumentVersion,
    Folder,
    MailRule,
    Project,
    ProjectMembership,
    Section,
)
from storage.domain.services import (
    DOCUMENT_VIEW_COLUMNS,
    CANONICAL_ZONE_NAMES,
    can_create_project,
    can_manage_project_participants,
    can_purge_trash,
    check_read_permission,
    copy_document,
    create_folder,
    create_project,
    create_section,
    delete_document,
    delete_documents,
    delete_folder,
    empty_trash,
    generate_share_link,
    get_accessible_documents,
    get_accessible_projects,
    get_accessible_sections,
    get_document_view,
    get_flat_folders,
    get_folder_documents,
    get_folder_tree,
    get_project_documents_for_view,
    get_project_documents_queryset,
    get_project_membership,
    get_sections_with_counts,
    get_trash,
    get_user_company,
    list_document_views,
    list_project_participants,
    move_document,
    move_document_to_section,
    move_documents,
    move_folder,
    normalize_document_view_config,
    rename_folder,
    restore_document,
    restore_folder,
    save_document_view,
    search_documents,
    update_project_member_role,
    upload_document,
)


RECENT_DOCUMENTS_SESSION_KEY = 'recent_documents'
RECENT_DOCUMENTS_LIMIT = 10

DOCUMENT_VIEW_COLUMN_LABELS = {
    'code': 'Шифр',
    'name': 'Название',
    'doc_type': 'Тип',
    'status': 'Статус',
    'zone': 'Зона',
    'version': 'Версия',
    'date': 'Дата',
    'author': 'Автор',
}

DEFAULT_COLUMN_WIDTHS = {
    'selection': 36,
    'code': 120,
    'name': 315,
    'doc_type': 100,
    'status': 100,
    'zone': 120,
    'version': 70,
    'date': 130,
    'author': 120,
}

DOCUMENT_TABLE_COLUMN_MAP = {
    'selection': 'selection',
    'code': 'code',
    'name': 'name',
    'doc_type': 'doc_type',
    'status': 'status',
    'zone': 'zone',
    'current_version': 'version',
    'updated_at': 'date',
    'created_by': 'author',
}

ALWAYS_VISIBLE_TABLE_COLUMNS = {'selection'}

FILTER_QUERY_KEYS = ('section', 'doc_type', 'status', 'code', 'name')

DOCUMENT_SORT_COLUMN_MAP = {
    'code': 'code',
    'name': 'name',
    'date': 'updated_at',
    'version': 'current_version',
}

HIGHLIGHT_ROW_SECONDS = 5


def _document_column_defs(config):
    visible = set(config['columns'])
    widths = config.get('widths') or {}
    return [
        {
            'key': key,
            'label': DOCUMENT_VIEW_COLUMN_LABELS[key],
            'visible': key in visible,
            'width': widths.get(key, DEFAULT_COLUMN_WIDTHS[key]),
        }
        for key in DOCUMENT_VIEW_COLUMNS
    ]


def _overlay_request_filters(request, base_config, folder=None):
    """Накладывает GET-параметры на сохранённый вид таблицы.

    GET-параметры ``section``, ``folder`` и фильтры в шапке таблицы
    применяются автоматически, без отдельной кнопки «Применить».
    """
    config = {
        'columns': base_config['columns'],
        'widths': base_config.get('widths', {}),
        'filters': dict(base_config['filters']),
        'sort': list(base_config['sort']),
    }
    requested = False

    for key in FILTER_QUERY_KEYS:
        if key not in request.GET:
            continue
        requested = True
        value = (request.GET.get(key) or '').strip()
        if value:
            config['filters'][key] = value
        else:
            config['filters'].pop(key, None)

    if 'sort' in request.GET:
        requested = True
        requested_sort = (request.GET.get('sort') or '').strip()
        if not requested_sort:
            config['sort'] = []
        elif requested_sort.lstrip('-') in DOCUMENT_SORT_COLUMN_MAP:
            config['sort'] = [requested_sort]

    if folder is not None:
        requested = True
        if folder.section_id:
            config['filters']['section'] = folder.section_id
        else:
            config['filters'].pop('section', None)

    if not requested:
        return base_config

    try:
        return normalize_document_view_config(config)
    except ValueError:
        # Некорректные GET-параметры не должны ронять страницу проекта.
        return base_config


def _folder_chain(folder):
    """Возвращает путь папки от корня до текущей."""
    chain = []
    current = folder
    seen = set()
    while current is not None and current.pk not in seen:
        chain.append(current)
        seen.add(current.pk)
        current = current.parent
    return list(reversed(chain))


def _project_breadcrumbs(project, section=None, folder=None):
    """Хлебные крошки страницы проекта: Проект > Раздел > Папка."""
    project_url = reverse('storage:project_detail', args=[project.pk])
    crumbs = [{
        'label': project.name,
        'url': project_url,
        'active': section is None and folder is None,
    }]

    if folder is not None:
        if folder.section_id:
            crumbs.append({
                'label': f'{folder.section.code} — {folder.section.name}',
                'url': f'{project_url}?section={folder.section_id}',
                'active': False,
            })
        chain = _folder_chain(folder)
        for index, node in enumerate(chain):
            crumbs.append({
                'label': node.name,
                'url': f'{project_url}?folder={node.pk}',
                'active': index == len(chain) - 1,
            })
    elif section is not None:
        crumbs.append({
            'label': f'{section.code} — {section.name}',
            'url': f'{project_url}?section={section.pk}',
            'active': True,
        })

    return crumbs


def _document_breadcrumbs(document):
    """Хлебные крошки документа: Проект > Раздел > Папка > Документ."""
    project_url = reverse('storage:project_detail', args=[document.project_id])
    crumbs = [{
        'label': document.project.name,
        'url': project_url,
        'active': False,
    }]
    section = document.section
    crumbs.append({
        'label': f'{section.code} — {section.name}',
        'url': f'{project_url}?section={section.pk}',
        'active': False,
    })
    for node in _folder_chain(document.folder):
        crumbs.append({
            'label': node.name,
            'url': f'{project_url}?folder={node.pk}',
            'active': False,
        })
    crumbs.append({
        'label': document.code,
        'url': '',
        'active': True,
    })
    return crumbs


def _remember_document(request, document_id):
    """Сохраняет в сессии только ID недавно открытых документов."""
    recent = [
        pk
        for pk in request.session.get(RECENT_DOCUMENTS_SESSION_KEY, [])
        if isinstance(pk, int) and pk != document_id
    ]
    request.session[RECENT_DOCUMENTS_SESSION_KEY] = [document_id, *recent][:RECENT_DOCUMENTS_LIMIT]


def _recent_documents(request, user, project):
    """Возвращает недавние документы, отфильтрованные по текущим правам."""
    ids = [
        pk
        for pk in request.session.get(RECENT_DOCUMENTS_SESSION_KEY, [])
        if isinstance(pk, int)
    ]
    if not ids:
        return []

    documents = {
        document.pk: document
        for document in get_accessible_documents(user).filter(
            project=project,
            pk__in=ids,
        ).select_related(
            'section',
            'current_version',
            'current_version__uploaded_by',
        )
    }
    return [documents[pk] for pk in ids if pk in documents][:RECENT_DOCUMENTS_LIMIT]


class ProjectListView(LoginRequiredMixin, ListView):
    model = Project
    template_name = 'storage/project_list.html'
    context_object_name = 'projects'
    ordering = ['code']

    def get_queryset(self):
        return get_accessible_projects(self.request.user)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        context['can_create_project'] = can_create_project(user)
        if context['can_create_project']:
            context['gip_candidates'] = (
                get_user_model()
                .objects.filter(
                    Q(userprofile__role__in=('gip', 'admin')) | Q(is_superuser=True)
                )
                .order_by('last_name', 'username')
            )
        return context


class ProjectDetailView(LoginRequiredMixin, CottonTableView):
    model = Document
    table_class = DocumentTable
    filterset_class = DocumentFilter
    paginate_by = 50
    title = 'Документы проекта'
    template_name = 'storage/project_detail.html'
    strict = False
    export_formats = ('xlsx',)
    export_name = 'documents'
    available_filter_buttons = []
    show_filter_line = False
    show_bulk_actions = False
    table_show_download_link = False
    table_empty_text = mark_safe(
        '<div class="doc-empty">'
        '<i class="bi bi-folder2-open doc-empty-icon" aria-hidden="true"></i>'
        '<div class="doc-empty-title">Документов пока нет</div>'
        '<div class="doc-empty-hint">'
        'Перетащите файл в таблицу или загрузите его кнопкой «Загрузить»'
        '</div>'
        '<label for="fileInput" class="btn btn-primary btn-sm">'
        '<i class="bi bi-plus-lg" aria-hidden="true"></i> '
        'Загрузить первый документ'
        '</label>'
        '</div>'
    )

    def _get_project(self):
        if not hasattr(self, '_project_cache'):
            self._project_cache = get_object_or_404(
                get_accessible_projects(self.request.user),
                pk=self.kwargs['pk'],
            )
        return self._project_cache

    def _get_selected_folder(self):
        if not hasattr(self, '_selected_folder_cache'):
            folder = None
            raw_folder = (self.request.GET.get('folder') or '').strip()
            if raw_folder:
                try:
                    folder_id = int(raw_folder)
                except (TypeError, ValueError):
                    folder_id = None
                if folder_id is not None:
                    folder = get_flat_folders(
                        self.request.user,
                        self._get_project(),
                    ).filter(pk=folder_id).select_related('section', 'parent').first()

            # Выбор раздела в панели фильтров отменяет фильтр папки,
            # если пользователь переключился на другой раздел.
            raw_section = (self.request.GET.get('section') or '').strip()
            if folder is not None and raw_section:
                try:
                    section_id = int(raw_section)
                except (TypeError, ValueError):
                    section_id = None
                if section_id is not None and section_id != folder.section_id:
                    folder = None
            self._selected_folder_cache = folder
        return self._selected_folder_cache

    def _get_selected_section(self):
        if not hasattr(self, '_selected_section_cache'):
            folder = self._get_selected_folder()
            if folder is not None and folder.section_id:
                self._selected_section_cache = folder.section
            else:
                section = None
                raw_section = (self.request.GET.get('section') or '').strip()
                if raw_section:
                    try:
                        section_id = int(raw_section)
                    except (TypeError, ValueError):
                        section_id = None
                    if section_id is not None:
                        section = get_accessible_sections(
                            self.request.user,
                            self._get_project(),
                        ).filter(pk=section_id).first()
                self._selected_section_cache = section
        return self._selected_section_cache

    def _get_view_state(self):
        if not hasattr(self, '_view_state_cache'):
            user = self.request.user
            project = self._get_project()
            document_views = list_document_views(user, project)
            selected_view = get_document_view(
                user,
                project,
                self.request.GET.get('view'),
            )

            if selected_view is not None:
                try:
                    base_config = normalize_document_view_config(selected_view.config)
                except ValueError:
                    # Повреждённый JSON в БД не должен ронять страницу проекта.
                    selected_view = None
                    base_config = normalize_document_view_config({})
            else:
                base_config = normalize_document_view_config({})

            view_config = _overlay_request_filters(
                self.request,
                base_config,
                folder=self._get_selected_folder(),
            )

            self._view_state_cache = (document_views, selected_view, view_config)
        return self._view_state_cache

    def get_queryset(self):
        return get_project_documents_queryset(
            self.request.user,
            self._get_project(),
            self._get_view_state()[2],
            folder=self._get_selected_folder(),
        )

    def get_table_data(self):
        if hasattr(self, 'object_list'):
            return self.object_list
        filterset = self.get_filterset(self.filterset_class)
        if filterset.is_valid():
            return filterset.qs
        return self.get_queryset()

    def get(self, request, *args, **kwargs):
        # Проверяем доступ к проекту до PaginationFixMixin, иначе его
        # обработчик Http404 превращает 404 в редирект.
        self._get_project()
        return super().get(request, *args, **kwargs)

    def get_filterset(self, filterset_class):
        filterset = super().get_filterset(filterset_class)
        section_field = filterset.form.fields.get('section')
        if section_field is not None:
            section_field.queryset = Section.objects.filter(
                project=self._get_project()
            )
        return filterset

    def get_filterset_kwargs(self, filterset_class):
        kwargs = super().get_filterset_kwargs(filterset_class)
        config = self._get_view_state()[2]
        data = self.request.GET.copy()
        changed = False
        for key in FILTER_QUERY_KEYS:
            value = config['filters'].get(key)
            if value and not data.get(key):
                data[key] = str(value)
                changed = True
        if changed:
            kwargs['data'] = data
        return kwargs

    def get_table_kwargs(self):
        kwargs = super().get_table_kwargs()
        visible = set(self._get_view_state()[2]['columns'])
        kwargs['exclude'] = [
            name
            for name in DocumentTable.base_columns
            if name not in ALWAYS_VISIBLE_TABLE_COLUMNS
            and DOCUMENT_TABLE_COLUMN_MAP.get(name, name) not in visible
        ]
        return kwargs

    def get_template_names(self):
        if getattr(self.request, 'htmx', False):
            return ['storage/_documents_panel.html']
        return [self.template_name]

    def get_table(self, **kwargs):
        table = super().get_table(**kwargs)
        config = self._get_view_state()[2]

        names = {column.name for column in table.columns}
        ordering = []
        for field in config['sort']:
            name = DOCUMENT_SORT_COLUMN_MAP[field.lstrip('-')]
            if name in names:
                prefix = '-' if field.startswith('-') else ''
                ordering.append(f'{prefix}{name}')
        if self.request.GET.get('apply') or not self.request.GET.get('sort'):
            table.order_by = ordering

        raw_highlight = (self.request.GET.get('highlight') or '').strip()
        table.highlight_id = int(raw_highlight) if raw_highlight.isdigit() else None

        return table

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        project = self._get_project()
        document_views, selected_view, view_config = self._get_view_state()

        can_manage_participants = can_manage_project_participants(user, project)
        active_tab = 'documents'
        requested_tab = self.request.GET.get('tab')
        if can_manage_participants and requested_tab in ('participants', 'trash'):
            active_tab = requested_tab

        sections = list(get_sections_with_counts(user, project))
        folder_roots = get_folder_tree(user, project)
        folders_by_section = {}
        project_folders = []
        for folder in folder_roots:
            if folder.section_id:
                folders_by_section.setdefault(folder.section_id, []).append(folder)
            else:
                project_folders.append(folder)
        for section in sections:
            section.folder_nodes = folders_by_section.get(section.pk, [])

        selected_folder = self._get_selected_folder()
        selected_section = self._get_selected_section()

        column_defs = _document_column_defs(view_config)
        documents = get_project_documents_for_view(
            user,
            project,
            view_config,
            folder=selected_folder,
        )
        active_search = (self.request.GET.get('q') or '').strip()
        if active_search:
            needle = active_search.lower()
            documents = [
                document
                for document in documents
                if needle in document['code'].lower()
                or needle in document['name'].lower()
            ]

        highlight_id = None
        raw_highlight = (self.request.GET.get('highlight') or '').strip()
        if raw_highlight.isdigit():
            highlight_id = int(raw_highlight)

        if selected_folder is not None:
            upload_target_name = selected_folder.name
        elif selected_section is not None:
            upload_target_name = f'раздел {selected_section.code}'
        else:
            upload_target_name = 'выбранный раздел'

        context.update({
            'project': project,
            'sections': sections,
            'folder_tree': folder_roots,
            'project_folders': project_folders,
            'all_folders': get_flat_folders(user, project),
            'recent_documents': _recent_documents(
                self.request,
                user,
                project,
            ),
            'documents': documents,
            'document_views': document_views,
            'selected_view': selected_view,
            'selected_folder': selected_folder,
            'selected_section': selected_section,
            'breadcrumbs': _project_breadcrumbs(
                project,
                section=selected_section,
                folder=selected_folder,
            ),
            'upload_target_name': upload_target_name,
            'highlight_id': highlight_id,
            'highlight_seconds': HIGHLIGHT_ROW_SECONDS,
            'active_search': active_search,
            'view_config': view_config,
            'widths_json': json.dumps(view_config['widths']),
            'active_filters': view_config['filters'],
            'column_defs': column_defs,
            'visible_columns': [column for column in column_defs if column['visible']],
            'active_sort': (
                self.request.GET.get('sort')
                or (view_config['sort'][0] if view_config['sort'] else '')
            ),
            'active_tab': active_tab,
            'can_manage_participants': can_manage_participants,
            'can_manage_structure': can_manage_participants,
            'can_purge_trash': can_purge_trash(user, project),
            'doc_type_choices': Document.DOC_TYPE_CHOICES,
            'status_choices': Document.STATUS_CHOICES,
            'CANONICAL_ZONE_NAMES': CANONICAL_ZONE_NAMES,
        })

        if can_manage_participants:
            context['participants'] = list_project_participants(user, project)
            context['membership_role_choices'] = ProjectMembership.ROLE_IN_PROJECT_CHOICES

        if active_tab == 'trash':
            context['trash'] = get_trash(user, project)

        return context


def _folder_archive_segments(folder, base_folder=None):
    """Возвращает относительные сегменты пути папки для ZIP."""
    segments = []
    current = folder
    while current is not None:
        if base_folder is not None and current.pk == base_folder.pk:
            break
        segments.append(get_valid_filename(current.name) or f'folder-{current.pk}')
        current = current.parent
    return list(reversed(segments))


def _zip_documents(documents, filename='documents.zip', base_folder=None):
    """Собирает ZIP из текущих версий документов с сохранением папок."""
    buffer = BytesIO()
    with zipfile.ZipFile(buffer, 'w', zipfile.ZIP_DEFLATED) as archive:
        for document in documents:
            version = document.current_version
            if version is None or not version.file:
                continue
            extension = version.file.name.rsplit('.', 1)[-1]
            if extension and len(extension) <= 8:
                extension = f'.{extension}'
            else:
                extension = ''
            document_filename = (
                get_valid_filename(document.code) or f'document-{document.pk}'
            )
            archive_name = '/'.join(
                _folder_archive_segments(document.folder, base_folder)
                + [f'{document_filename}_v{version.version_number}{extension}']
            )
            try:
                with version.file.open('rb') as source:
                    archive.writestr(archive_name, source.read())
            except (FileNotFoundError, OSError):
                continue

    response = HttpResponse(buffer.getvalue(), content_type='application/zip')
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


class ProjectCreateView(LoginRequiredMixin, View):
    def post(self, request):
        if not can_create_project(request.user):
            return HttpResponseForbidden('Нет прав на создание проекта')

        company = get_user_company(request.user)
        if company is None and request.user.is_superuser:
            company = Company.objects.first()
        if company is None:
            messages.error(request, 'Не удалось определить компанию пользователя.')
            return redirect('storage:project_list')

        gip = None
        gip_id = request.POST.get('gip')
        if gip_id:
            gip = get_user_model().objects.filter(pk=gip_id).first()

        try:
            project = create_project(
                company=company,
                code=request.POST.get('code', '').strip(),
                name=request.POST.get('name', '').strip(),
                customer=request.POST.get('customer', '').strip(),
                gip=gip,
                created_by=request.user,
            )
        except (PermissionDenied, ValueError) as exc:
            messages.error(request, str(exc))
            return redirect('storage:project_list')

        messages.success(request, f'Проект «{project.name}» создан.')
        return redirect('storage:project_detail', pk=project.pk)


class SectionCreateView(LoginRequiredMixin, View):
    def post(self, request, project_id):
        project = get_object_or_404(
            get_accessible_projects(request.user),
            pk=project_id,
        )
        try:
            create_section(
                project=project,
                code=request.POST.get('code', ''),
                name=request.POST.get('name', ''),
                order=request.POST.get('order') or 0,
                user=request.user,
            )
        except (PermissionDenied, ValueError) as exc:
            messages.error(request, str(exc))
        else:
            messages.success(request, 'Раздел создан.')
        return redirect('storage:project_detail', pk=project.pk)


class FolderCreateView(LoginRequiredMixin, View):
    def post(self, request, project_id):
        project = get_object_or_404(
            get_accessible_projects(request.user),
            pk=project_id,
        )
        section = None
        section_id = request.POST.get('section')
        if section_id:
            section = Section.objects.filter(project=project, pk=section_id).first()
            if section is None:
                messages.error(request, 'Раздел не найден.')
                return redirect('storage:project_detail', pk=project.pk)
        parent = None
        parent_id = request.POST.get('parent')
        if parent_id:
            parent = Folder.objects.filter(project=project, pk=parent_id).first()
            if parent is None:
                messages.error(request, 'Родительская папка не найдена.')
                return redirect('storage:project_detail', pk=project.pk)

        try:
            create_folder(
                user=request.user,
                project=project,
                name=request.POST.get('name', ''),
                section=section,
                parent=parent,
            )
        except (PermissionDenied, ValueError) as exc:
            messages.error(request, str(exc))
        else:
            messages.success(request, 'Папка создана.')
        return redirect('storage:project_detail', pk=project.pk)


class FolderCreateFormView(LoginRequiredMixin, TemplateView):
    """HTML-фрагмент формы создания папки для Bootstrap modal (HTMX)."""

    template_name = 'storage/_folder_create_form.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        project = get_object_or_404(
            get_accessible_projects(self.request.user),
            pk=self.kwargs['project_id'],
        )
        if not can_manage_project_participants(self.request.user, project):
            raise PermissionDenied("Недостаточно прав для создания папки")

        parent = None
        parent_id = (self.request.GET.get('parent') or '').strip()
        if parent_id.isdigit():
            parent = Folder.objects.filter(
                project=project,
                pk=int(parent_id),
                is_deleted=False,
            ).select_related('section').first()

        section = parent.section if parent is not None else None
        section_id = (self.request.GET.get('section') or '').strip()
        if section is None and section_id.isdigit():
            section = Section.objects.filter(
                project=project,
                pk=int(section_id),
            ).first()

        context.update({
            'project': project,
            'parent': parent,
            'section': section,
        })
        return context


class FolderRenameFormView(LoginRequiredMixin, TemplateView):
    """HTML-фрагмент формы переименования папки для Bootstrap modal (HTMX)."""

    template_name = 'storage/_folder_rename_form.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        folder = get_object_or_404(
            Folder.objects.select_related('project', 'section', 'parent'),
            pk=self.kwargs['folder_id'],
            project__in=get_accessible_projects(self.request.user),
        )
        if not can_manage_project_participants(self.request.user, folder.project):
            raise PermissionDenied("Недостаточно прав для переименования папки")
        context.update({'folder': folder, 'project': folder.project})
        return context


class FolderRenameView(LoginRequiredMixin, View):
    def post(self, request, folder_id):
        folder = get_object_or_404(
            Folder,
            pk=folder_id,
            project__in=get_accessible_projects(request.user),
        )
        try:
            rename_folder(request.user, folder, request.POST.get('name', ''))
        except (PermissionDenied, ValueError) as exc:
            messages.error(request, str(exc))
        else:
            messages.success(request, 'Папка переименована.')
        return redirect('storage:project_detail', pk=folder.project_id)


class FolderDeleteView(LoginRequiredMixin, View):
    def post(self, request, folder_id):
        folder = get_object_or_404(
            Folder,
            pk=folder_id,
            project__in=get_accessible_projects(request.user),
        )
        try:
            delete_folder(request.user, folder)
        except (PermissionDenied, ValueError) as exc:
            messages.error(request, str(exc))
        else:
            messages.success(request, 'Папка перемещена в корзину.')
        return redirect('storage:project_detail', pk=folder.project_id)


class FolderDownloadView(LoginRequiredMixin, View):
    def get(self, request, folder_id):
        folder = get_object_or_404(
            Folder,
            pk=folder_id,
            project__in=get_accessible_projects(request.user),
        )
        try:
            documents = list(get_folder_documents(request.user, folder))
        except (PermissionDenied, ValueError) as exc:
            return HttpResponseForbidden(str(exc))
        return _zip_documents(documents, 'folder.zip', base_folder=folder)


class DocumentShareLinkView(LoginRequiredMixin, View):
    def get(self, request, document_id):
        document = get_object_or_404(
            get_accessible_documents(request.user),
            pk=document_id,
        )
        try:
            token = generate_share_link(request.user, document)
        except (PermissionDenied, ValueError) as exc:
            return JsonResponse({'error': str(exc)}, status=403)
        url = request.build_absolute_uri(
            reverse('storage:document_detail', args=[document.pk])
        )
        return JsonResponse({'token': token, 'url': f'{url}?share={token}'})


class DocumentMoveView(LoginRequiredMixin, View):
    """AJAX-перемещение документа при drag-and-drop на папку или раздел."""

    def post(self, request, document_id):
        document = get_object_or_404(
            get_accessible_documents(request.user),
            pk=document_id,
            is_deleted=False,
        )
        folder_id = (request.POST.get('folder_id') or '').strip()
        section_id = (request.POST.get('section_id') or '').strip()

        try:
            if folder_id:
                folder = Folder.objects.filter(
                    project=document.project,
                    pk=folder_id,
                    is_deleted=False,
                ).first()
                if folder is None:
                    raise ValueError('Папка не найдена.')
                move_document(request.user, document, folder)
            elif section_id:
                section = get_object_or_404(
                    get_accessible_sections(request.user, document.project),
                    pk=section_id,
                )
                move_document_to_section(request.user, document, section)
            else:
                move_document(request.user, document, None)
        except PermissionDenied as exc:
            return JsonResponse(
                {'status': 'error', 'error': str(exc)},
                status=403,
            )
        except ValueError as exc:
            return JsonResponse(
                {'status': 'error', 'error': str(exc)},
                status=400,
            )

        return JsonResponse({
            'status': 'ok',
            'document_id': document.pk,
            'folder_id': document.folder_id,
            'section_id': document.section_id,
        })


class DocumentDeleteView(LoginRequiredMixin, View):
    """AJAX-отправка документа в корзину из контекстного меню."""

    def post(self, request, document_id):
        document = get_object_or_404(
            get_accessible_documents(request.user),
            pk=document_id,
            is_deleted=False,
        )
        try:
            delete_document(request.user, document)
        except PermissionDenied as exc:
            return JsonResponse(
                {'status': 'error', 'error': str(exc)},
                status=403,
            )
        except ValueError as exc:
            return JsonResponse(
                {'status': 'error', 'error': str(exc)},
                status=400,
            )
        return JsonResponse({
            'status': 'ok',
            'document_id': document.pk,
        })


class DocumentCopyView(LoginRequiredMixin, View):
    """AJAX-копирование документа при drop с зажатым Ctrl."""

    def post(self, request, document_id):
        document = get_object_or_404(
            get_accessible_documents(request.user),
            pk=document_id,
            is_deleted=False,
        )
        folder_id = (request.POST.get('folder_id') or '').strip()
        section_id = (request.POST.get('section_id') or '').strip()
        folder = None
        section = None

        try:
            if folder_id:
                folder = Folder.objects.filter(
                    project=document.project,
                    pk=folder_id,
                    is_deleted=False,
                ).first()
                if folder is None:
                    raise ValueError('Папка не найдена.')
            elif section_id:
                section = get_object_or_404(
                    get_accessible_sections(request.user, document.project),
                    pk=section_id,
                )
            else:
                folder = document.folder
                section = document.section

            copied = copy_document(
                request.user,
                document,
                folder,
                section,
                LocalFileStorage(),
            )
        except PermissionDenied as exc:
            return JsonResponse(
                {'status': 'error', 'error': str(exc)},
                status=403,
            )
        except ValueError as exc:
            return JsonResponse(
                {'status': 'error', 'error': str(exc)},
                status=400,
            )

        return JsonResponse({
            'status': 'ok',
            'document_id': copied.pk,
            'code': copied.code,
            'folder_id': copied.folder_id,
            'section_id': copied.section_id,
        })


class FolderMoveView(LoginRequiredMixin, View):
    """AJAX-перемещение папки в другую папку или раздел."""

    def post(self, request, folder_id):
        folder = get_object_or_404(
            Folder.objects.select_related('project', 'section', 'parent'),
            pk=folder_id,
            project__in=get_accessible_projects(request.user),
            is_deleted=False,
        )
        parent_id = (request.POST.get('parent_id') or '').strip()
        section_id = (request.POST.get('section_id') or '').strip()
        parent = None
        section = None

        try:
            if parent_id:
                parent = Folder.objects.filter(
                    project=folder.project,
                    pk=parent_id,
                    is_deleted=False,
                ).first()
                if parent is None:
                    raise ValueError('Целевая папка не найдена.')
            elif section_id:
                section = Section.objects.filter(
                    project=folder.project,
                    pk=section_id,
                ).first()
                if section is None:
                    raise ValueError('Целевой раздел не найден.')

            move_folder(request.user, folder, parent=parent, section=section)
        except PermissionDenied as exc:
            return JsonResponse(
                {'status': 'error', 'error': str(exc)},
                status=403,
            )
        except ValueError as exc:
            return JsonResponse(
                {'status': 'error', 'error': str(exc)},
                status=400,
            )

        return JsonResponse({
            'status': 'ok',
            'folder_id': folder.pk,
            'parent_id': folder.parent_id,
            'section_id': folder.section_id,
        })


class DocumentBulkActionView(LoginRequiredMixin, View):
    def post(self, request, project_id):
        project = get_object_or_404(
            get_accessible_projects(request.user),
            pk=project_id,
        )
        document_ids = request.POST.getlist('document_ids')
        documents = list(
            get_accessible_documents(request.user).filter(
                project=project,
                pk__in=document_ids,
                is_deleted=False,
            )
        )
        if not documents:
            messages.error(request, 'Не выбраны документы.')
            return redirect('storage:project_detail', pk=project.pk)

        action = request.POST.get('action')
        try:
            if action == 'download':
                return _zip_documents(documents)
            if action == 'move':
                folder = None
                folder_id = request.POST.get('folder')
                if folder_id:
                    folder = Folder.objects.filter(
                        project=project,
                        pk=folder_id,
                        is_deleted=False,
                    ).first()
                    if folder is None:
                        raise ValueError('Папка не найдена.')
                move_documents(request.user, documents, folder)
                messages.success(request, f'Перемещено документов: {len(documents)}.')
            elif action == 'delete':
                delete_documents(request.user, documents)
                messages.success(request, f'В корзину отправлено: {len(documents)}.')
            else:
                messages.error(request, 'Неизвестное действие.')
        except (PermissionDenied, ValueError) as exc:
            messages.error(request, str(exc))

        return redirect('storage:project_detail', pk=project.pk)


class TrashRestoreView(LoginRequiredMixin, View):
    def post(self, request, project_id):
        project = get_object_or_404(
            get_accessible_projects(request.user),
            pk=project_id,
        )
        item_type = request.POST.get('item_type')
        item_id = request.POST.get('item_id')
        try:
            if item_type == 'document':
                document = get_object_or_404(
                    Document,
                    pk=item_id,
                    project=project,
                    is_deleted=True,
                )
                restore_document(request.user, document)
            elif item_type == 'folder':
                folder = get_object_or_404(
                    Folder,
                    pk=item_id,
                    project=project,
                    is_deleted=True,
                )
                restore_folder(request.user, folder)
            else:
                raise ValueError('Неизвестный тип элемента корзины')
        except (PermissionDenied, ValueError) as exc:
            messages.error(request, str(exc))
        else:
            messages.success(request, 'Элемент восстановлен.')
        return redirect(f"{reverse('storage:project_detail', args=[project.pk])}?tab=trash")


class TrashClearView(LoginRequiredMixin, View):
    """Безвозвратная очистка корзины проекта (только администратор)."""

    def post(self, request, project_id):
        project = get_object_or_404(
            get_accessible_projects(request.user),
            pk=project_id,
        )
        try:
            result = empty_trash(request.user, project, LocalFileStorage())
        except PermissionDenied as exc:
            return HttpResponseForbidden(str(exc))
        except ValueError as exc:
            messages.error(request, str(exc))
        else:
            messages.success(
                request,
                'Корзина очищена: удалено документов — '
                f"{result['documents']}, папок — {result['folders']}.",
            )
        return redirect(f"{reverse('storage:project_detail', args=[project.pk])}?tab=trash")


class DocumentViewSaveView(LoginRequiredMixin, View):
    def post(self, request, project_id):
        project = get_object_or_404(
            get_accessible_projects(request.user),
            pk=project_id,
        )

        raw_widths = (request.POST.get('widths') or '').strip()
        widths = {}
        if raw_widths:
            try:
                widths = json.loads(raw_widths)
            except (TypeError, ValueError):
                messages.error(request, 'Некорректные ширины столбцов.')
                return redirect('storage:project_detail', pk=project.pk)

        config = {
            'columns': request.POST.getlist('columns'),
            'widths': widths,
            'filters': {
                'section': request.POST.get('section', ''),
                'doc_type': request.POST.get('doc_type', ''),
                'status': request.POST.get('status', ''),
            },
            'sort': request.POST.get('sort', ''),
        }

        try:
            view = save_document_view(
                user=request.user,
                project=project,
                name=request.POST.get('name', ''),
                config=config,
                is_shared=request.POST.get('is_shared') == 'on',
            )
        except PermissionDenied as exc:
            return HttpResponseForbidden(str(exc))
        except ValueError as exc:
            messages.error(request, str(exc))
            return redirect('storage:project_detail', pk=project.pk)

        messages.success(request, f'Вид «{view.name}» сохранён.')
        project_url = reverse('storage:project_detail', args=[project.pk])
        return redirect(f'{project_url}?view={view.pk}')


class MembershipRoleUpdateView(LoginRequiredMixin, View):
    def post(self, request, project_id, membership_id):
        project = get_object_or_404(
            get_accessible_projects(request.user),
            pk=project_id,
        )
        try:
            membership = get_project_membership(
                request.user,
                project,
                membership_id,
            )
            if membership is None:
                raise Http404("Участник проекта не найден")

            update_project_member_role(
                request.user,
                membership,
                request.POST.get('role', ''),
            )
        except PermissionDenied as exc:
            return HttpResponseForbidden(str(exc))
        except ValueError as exc:
            messages.error(request, str(exc))
            project_url = reverse('storage:project_detail', args=[project.pk])
            return redirect(f'{project_url}?tab=participants')

        messages.success(
            request,
            f'Роль участника {membership.user.get_full_name() or membership.user.username} обновлена.',
        )
        project_url = reverse('storage:project_detail', args=[project.pk])
        return redirect(f'{project_url}?tab=participants')


class SearchView(LoginRequiredMixin, TemplateView):
    template_name = 'storage/search.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        user = self.request.user
        project = get_object_or_404(
            get_accessible_projects(user),
            pk=self.kwargs['project_id'],
        )
        sections = get_accessible_sections(user, project)

        query = (self.request.GET.get('q') or '').strip()

        selected_section = None
        section_id = (self.request.GET.get('section') or '').strip()
        if section_id:
            try:
                selected_section = sections.filter(pk=int(section_id)).first()
            except (TypeError, ValueError):
                selected_section = None

        doc_type_choices = Document.DOC_TYPE_CHOICES
        doc_type = (self.request.GET.get('doc_type') or '').strip()
        if doc_type not in dict(doc_type_choices):
            doc_type = ''

        status_choices = Document.STATUS_CHOICES
        status = (self.request.GET.get('status') or '').strip()
        if status not in dict(status_choices):
            status = ''

        documents = search_documents(
            user=user,
            project=project,
            query=query,
            section=selected_section,
            doc_type=doc_type or None,
            status=status or None,
        )

        context.update({
            'project': project,
            'documents': documents,
            'query': query,
            'sections': sections,
            'selected_section': selected_section,
            'selected_doc_type': doc_type,
            'selected_status': status,
            'doc_type_choices': doc_type_choices,
            'status_choices': status_choices,
        })
        return context


class DocumentDetailView(LoginRequiredMixin, DetailView):
    model = Document
    template_name = 'storage/document_detail.html'
    context_object_name = 'document'

    def dispatch(self, request, *args, **kwargs):
        share_token = request.GET.get('share')
        if share_token:
            try:
                token = uuid.UUID(str(share_token))
            except (TypeError, ValueError, AttributeError):
                raise Http404
            document = (
                Document.objects.filter(
                    pk=kwargs['pk'],
                    public_link_token=token,
                    is_deleted=False,
                )
                .select_related(
                    'project',
                    'section',
                    'current_version',
                    'current_version__uploaded_by',
                )
                .first()
            )
            if document is None:
                raise Http404
            self.object = document
            _remember_document(request, document.pk)
            context = self.get_context_data(object=document)
            context['public_view'] = True
            return self.render_to_response(context)
        return super().dispatch(request, *args, **kwargs)

    def get_queryset(self):
        return get_accessible_documents(self.request.user).select_related(
            'project',
            'section',
            'current_version',
            'current_version__uploaded_by',
        )

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        if not check_read_permission(request.user, self.object):
            return HttpResponseForbidden('Нет прав на просмотр документа')
        _remember_document(request, self.object.pk)
        context = self.get_context_data(object=self.object)
        return self.render_to_response(context)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['breadcrumbs'] = _document_breadcrumbs(self.object)
        context['versions'] = self.object.versions.select_related('uploaded_by').order_by('-version_number')
        context['issues'] = self.object.issues.select_related('author', 'assignee', 'version').order_by('-created_at')
        return context


class DocumentUploadAjaxView(LoginRequiredMixin, View):
    def handle_no_permission(self):
        return JsonResponse({'error': 'Требуется авторизация'}, status=403)

    def post(self, request, project_id):
        project = get_accessible_projects(request.user).filter(pk=project_id).first()
        if project is None:
            return JsonResponse({'error': 'Проект не найден'}, status=404)

        section_id = request.POST.get('section_id')
        folder_id = request.POST.get('folder_id')
        file = request.FILES.get('file')
        if file is None:
            return JsonResponse({'error': 'Файл не выбран'}, status=400)

        folder = None
        if folder_id:
            folder = Folder.objects.filter(
                project=project,
                pk=folder_id,
                is_deleted=False,
            ).first()
            if folder is None:
                return JsonResponse({'error': 'Папка не найдена'}, status=400)

        if folder is not None and folder.section_id is not None:
            section = folder.section
            if section_id and str(section.pk) != str(section_id):
                return JsonResponse(
                    {'error': 'Раздел не совпадает с разделом папки'},
                    status=400,
                )
        else:
            if not section_id:
                return JsonResponse({'error': 'Не выбран раздел'}, status=400)
            try:
                section = Section.objects.get(
                    pk=section_id,
                    project__company=project.company,
                )
            except (Section.DoesNotExist, ValueError, TypeError):
                return JsonResponse({'error': 'Раздел не найден'}, status=400)

        if folder is not None and folder.project_id != project.pk:
            return JsonResponse({'error': 'Папка не найдена'}, status=400)

        code = (request.POST.get('code') or file.name).strip()
        name = (request.POST.get('name') or file.name).strip()
        doc_type = request.POST.get('doc_type') or 'other'

        try:
            version = upload_document(
                user=request.user,
                project=project,
                section=section,
                code=code,
                name=name,
                doc_type=doc_type,
                file=file,
                comment=request.POST.get('comment', ''),
                folder=folder,
            )
        except (PermissionDenied, ValueError) as exc:
            return JsonResponse({'error': str(exc)}, status=400)

        response = JsonResponse({
            'status': 'ok',
            'version': version.version_number,
        })
        response['X-Document-Id'] = str(version.document_id)
        return response


class DocumentDownloadView(LoginRequiredMixin, View):
    def get(self, request, version_id):
        versions = DocumentVersion.objects.select_related('document__project').filter(
            document__is_deleted=False,
        )

        if not request.user.is_superuser:
            company = get_user_company(request.user)
            if company is None:
                return HttpResponseForbidden('Нет прав на скачивание файла')
            versions = versions.filter(document__project__company=company)

        version = get_object_or_404(versions, pk=version_id)

        if not check_read_permission(request.user, version.document):
            return HttpResponseForbidden('Нет прав на скачивание файла')

        _remember_document(request, version.document_id)

        if not version.file:
            return HttpResponseNotFound('Файл не найден')

        try:
            file_handle = version.file.open('rb')
        except FileNotFoundError:
            return HttpResponseNotFound('Файл не найден')

        filename = version.file.name.rsplit('/', 1)[-1]
        return FileResponse(file_handle, as_attachment=True, filename=filename)


class ExportVedomostView(LoginRequiredMixin, View):
    def get(self, request, project_id):
        project = get_object_or_404(get_accessible_projects(request.user), pk=project_id)

        documents = get_accessible_documents(request.user).filter(project=project).select_related(
            'section',
            'current_version',
            'current_version__uploaded_by',
        ).order_by('section__order', 'code')

        section_id = request.GET.get('section_id')
        if section_id:
            try:
                documents = documents.filter(section_id=int(section_id))
            except (TypeError, ValueError):
                return HttpResponseBadRequest('Неверный идентификатор раздела')

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Ведомость'

        sheet.append(['Шифр', 'Название', 'Тип', 'Версия', 'Дата', 'Автор'])
        for cell in sheet[1]:
            cell.font = Font(bold=True)

        for row, document in enumerate(documents, start=2):
            version = document.current_version

            sheet.cell(row=row, column=1, value=document.code)

            name_cell = sheet.cell(row=row, column=2, value=document.name)
            if version is not None:
                download_url = request.build_absolute_uri(
                    reverse('storage:document_download', args=[version.pk])
                )
                name_cell.hyperlink = download_url
                name_cell.font = Font(color='0563C1', underline='single')

            sheet.cell(row=row, column=3, value=document.get_doc_type_display())
            sheet.cell(row=row, column=4, value=version.version_number if version else '-')
            sheet.cell(
                row=row,
                column=5,
                value=self._format_datetime(version.uploaded_at if version else None),
            )
            sheet.cell(row=row, column=6, value=self._format_author(version))

        for column, width in {'A': 18, 'B': 40, 'C': 22, 'D': 10, 'E': 18, 'F': 24}.items():
            sheet.column_dimensions[column].width = width

        buffer = BytesIO()
        workbook.save(buffer)
        buffer.seek(0)

        return FileResponse(
            buffer,
            as_attachment=True,
            filename='vedomost.xlsx',
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )

    @staticmethod
    def _format_datetime(value):
        if value is None:
            return ''
        if timezone.is_aware(value):
            value = timezone.localtime(value)
        return value.strftime('%d.%m.%Y %H:%M')

    @staticmethod
    def _format_author(version):
        if version is None:
            return ''
        return version.uploaded_by.get_full_name() or version.uploaded_by.username


class SimulateMailImportView(LoginRequiredMixin, View):
    def post(self, request, project_id):
        project = get_object_or_404(get_accessible_projects(request.user), pk=project_id)

        rule = MailRule.objects.filter(project=project).order_by('pk').first()
        if rule is None:
            messages.error(request, 'Для проекта не настроены правила импорта из почты.')
            return redirect('storage:project_detail', pk=project.pk)

        filename = 'Исходные_данные_из_почты.pdf'
        fake_file = SimpleUploadedFile(
            filename,
            b'%PDF-1.4 simulated mail attachment',
            content_type='application/pdf',
        )

        try:
            upload_document(
                user=request.user,
                project=project,
                section=rule.section,
                code='ИЗ-ПОЧТЫ-ИД',
                name='Исходные данные из почты',
                doc_type=rule.doc_type,
                file=fake_file,
                comment='Импортировано из почты (симуляция)',
            )
        except (PermissionDenied, ValueError) as exc:
            messages.error(request, str(exc))
            return redirect('storage:project_detail', pk=project.pk)

        messages.success(request, 'Письмо обработано, файл добавлен.')
        return redirect('storage:project_detail', pk=project.pk)
